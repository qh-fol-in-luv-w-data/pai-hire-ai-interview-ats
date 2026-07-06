"""Generate PAI HR non-technical overview document"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def heading(doc, text, level, color='131B2E'):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor.from_string(color)
    return p

def body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    return p

def bullet(doc, icon, bold_text, desc):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(icon + '  ')
    r1.font.size = Pt(12)
    r2 = p.add_run(bold_text)
    r2.bold = True
    if desc:
        p.add_run('  —  ' + desc)
    return p

doc = Document()
for section in doc.sections:
    section.top_margin    = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin   = Cm(3.0)
    section.right_margin  = Cm(3.0)

doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11.5)

# ── COVER ──────────────────────────────────────────────────────
doc.add_paragraph()
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('PAI HR')
r.bold = True; r.font.size = Pt(32)
r.font.color.rgb = RGBColor(0x00, 0x51, 0xD5)

t2 = doc.add_paragraph()
t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = t2.add_run('Hệ Thống Tuyển Dụng Thông Minh')
r2.font.size = Pt(16)
r2.font.color.rgb = RGBColor(0x37, 0x41, 0x51)

doc.add_paragraph()
t3 = doc.add_paragraph()
t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r3 = t3.add_run('Tài liệu giới thiệu dành cho HR & Ban lãnh đạo')
r3.font.size = Pt(12)
r3.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
r3.italic = True

doc.add_paragraph()
sep = doc.add_paragraph()
sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
sep.add_run('─' * 40).font.color.rgb = RGBColor(0xE5, 0xE7, 0xEB)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# 1. PAI HR LÀ GÌ?
# ══════════════════════════════════════════════════════════════
heading(doc, '1. PAI HR là gì?', 1, '0051D5')

body(doc,
    'PAI HR là hệ thống tuyển dụng tự động, giúp đội ngũ HR xử lý toàn bộ quy trình '
    'từ lúc ứng viên nộp hồ sơ cho đến khi hoàn thành vòng phỏng vấn — '
    'mà không cần sắp xếp lịch hay xem xét thủ công từng CV một.'
)

body(doc,
    'Thay vì HR phải đọc hàng chục CV mỗi ngày, hệ thống sẽ tự động chấm điểm, '
    'gửi email thông báo và mời ứng viên đạt yêu cầu vào vòng phỏng vấn AI — '
    'tất cả diễn ra trong vòng vài phút sau khi ứng viên nộp hồ sơ.'
)

heading(doc, '1.1 Lợi ích chính', 2, '374151')
benefits = [
    ('⏱', 'Tiết kiệm thời gian', 'CV được chấm điểm tự động, HR chỉ xem xét hồ sơ đã qua sàng lọc.'),
    ('⚖', 'Công bằng & nhất quán', 'Mọi ứng viên được đánh giá theo cùng một tiêu chí, không bị ảnh hưởng bởi cảm tính.'),
    ('📱', 'Ứng viên không cần cài app', 'Toàn bộ phỏng vấn diễn ra trên trình duyệt, chỉ cần bấm link trong email.'),
    ('📊', 'Dashboard tổng quan', 'HR theo dõi pipeline tuyển dụng theo thời gian thực — bao nhiêu hồ sơ, bao nhiêu đạt.'),
    ('✉', 'Email tự động', 'Hệ thống tự gửi thư mời phỏng vấn hoặc thư từ chối, không cần làm thủ công.'),
]
for icon, bold, desc in benefits:
    bullet(doc, icon, bold, desc)

# ══════════════════════════════════════════════════════════════
# 2. QUY TRÌNH TUYỂN DỤNG
# ══════════════════════════════════════════════════════════════
heading(doc, '2. Quy Trình Tuyển Dụng', 1, '0051D5')

body(doc, 'Toàn bộ quy trình gồm 6 bước, từ khi ứng viên thấy tin tuyển dụng '
         'đến khi HR nhận được kết quả phỏng vấn để đưa ra quyết định cuối cùng.')

steps = [
    ('BƯỚC 1', '📋 Ứng viên xem tin tuyển dụng',
     'Ứng viên truy cập trang web của PAI HR, xem danh sách các vị trí đang tuyển, '
     'đọc mô tả công việc và yêu cầu từng vị trí.'),
    ('BƯỚC 2', '📎 Nộp hồ sơ',
     'Ứng viên điền thông tin cơ bản (họ tên, email, số điện thoại) và tải lên CV. '
     'Hệ thống xác nhận nhận hồ sơ ngay lập tức và cấp mã hồ sơ riêng.'),
    ('BƯỚC 3', '🤖 Hệ thống AI chấm điểm CV',
     'Ngay sau khi nhận hồ sơ, hệ thống AI tự động đọc và chấm điểm CV theo 10 tiêu chí '
     '(kinh nghiệm, học vấn, kỹ năng, dự án…). Quá trình này mất khoảng 1–2 phút. '
     'Điểm tối đa là 10, ngưỡng đạt là 7.5 điểm.'),
    ('BƯỚC 4A', '✅ Email mời phỏng vấn (nếu đạt)',
     'Ứng viên đạt ≥ 7.5 điểm sẽ nhận email mời tham gia phỏng vấn AI trực tuyến. '
     'Email có nút bấm trực tiếp, ứng viên chỉ cần click là vào ngay trang phỏng vấn. '
     'Không cần đặt lịch, không cần chờ HR liên lạc.'),
    ('BƯỚC 4B', '❌ Email thông báo (nếu không đạt)',
     'Ứng viên không đạt ngưỡng sẽ nhận email lịch sự cảm ơn và khuyến khích '
     'ứng tuyển các vị trí phù hợp hơn trong tương lai. Điểm số không được tiết lộ.'),
    ('BƯỚC 5', '🎙 Phỏng vấn AI trực tuyến',
     'Ứng viên bấm link trong email, vào thẳng trang phỏng vấn. Hệ thống đọc to từng câu hỏi '
     'bằng giọng AI, ứng viên ghi âm câu trả lời ngay trên điện thoại hoặc máy tính. '
     'Phỏng vấn gồm 6 câu hỏi chia 3 phần, thường mất khoảng 10–15 phút.'),
    ('BƯỚC 6', '👩‍💼 HR xem xét & quyết định',
     'Sau khi ứng viên nộp bài, HR đăng nhập vào dashboard, nghe lại câu trả lời, '
     'chấm điểm từng câu và đưa ra quyết định tuyển dụng cuối cùng.'),
]

for label, title, desc in steps:
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0)

    # Label badge
    r_label = p.add_run(f'  {label}  ')
    r_label.bold = True
    r_label.font.size = Pt(9)
    r_label.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Title
    r_title = p.add_run(f'  {title}')
    r_title.bold = True
    r_title.font.size = Pt(12)
    r_title.font.color.rgb = RGBColor(0x13, 0x1B, 0x2E)

    # Description
    p2 = doc.add_paragraph(desc)
    p2.paragraph_format.left_indent = Cm(1.0)
    p2.paragraph_format.space_after = Pt(2)
    for run in p2.runs:
        run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
        run.font.size = Pt(11)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# 3. PHỎNG VẤN AI — HOẠT ĐỘNG NHƯ THẾ NÀO?
# ══════════════════════════════════════════════════════════════
heading(doc, '3. Phỏng Vấn AI — Hoạt Động Như Thế Nào?', 1, '0051D5')

body(doc,
    'Buổi phỏng vấn AI hoàn toàn tự động, không cần HR có mặt. '
    'Ứng viên nghe câu hỏi và ghi âm câu trả lời — giống như một cuộc phỏng vấn bình thường, '
    'chỉ khác là phía bên kia là giọng AI thay vì người thật.'
)

heading(doc, '3.1 Cấu trúc 3 phần — 6 câu hỏi', 2, '374151')

parts = [
    ('Phần 1', 'Câu 01 & 02', 'Chuyên Môn Kỹ Thuật',
     'Hỏi về kiến thức chuyên ngành, công nghệ ứng viên sử dụng, '
     'và cách xử lý các tình huống kỹ thuật thực tế trong công việc.'),
    ('Phần 2', 'Câu 03 & 04', 'Kỹ Năng Mềm',
     'Hỏi về cách ứng viên giao tiếp, làm việc nhóm, '
     'xử lý áp lực deadline, và giải quyết xung đột với đồng nghiệp.'),
    ('Phần 3', 'Câu 05 & 06', 'Kinh Nghiệm Thực Tế',
     'Hỏi về các dự án đã tham gia, vai trò cụ thể, kết quả đạt được, '
     'và bài học rút ra từ những thất bại trong quá khứ.'),
]

t = doc.add_table(rows=1, cols=4)
t.style = 'Table Grid'
for i, h in enumerate(['', 'Câu hỏi', 'Chủ đề', 'Nội dung đánh giá']):
    t.rows[0].cells[i].text = h
    t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    set_cell_bg(t.rows[0].cells[i], '131B2E')
    t.rows[0].cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF,0xFF,0xFF)

bgs = ['EFF6FF', 'ECFDF5', 'FFFBEB']
for i, (phan, cau, chu_de, noi_dung) in enumerate(parts):
    row = t.add_row()
    row.cells[0].text = phan
    row.cells[0].paragraphs[0].runs[0].bold = True
    row.cells[1].text = cau
    row.cells[2].text = chu_de
    row.cells[2].paragraphs[0].runs[0].bold = True
    row.cells[3].text = noi_dung
    for c in row.cells: set_cell_bg(c, bgs[i])

doc.add_paragraph()

heading(doc, '3.2 Câu hỏi mẫu theo từng phần', 2, '374151')

sample_parts = [
    ('Phần 1 — Chuyên Môn Kỹ Thuật', '1D4ED8', [
        'Bạn hãy mô tả quy trình làm việc thông thường của bạn khi bắt đầu một dự án mới từ đầu.',
        'Kể về một vấn đề kỹ thuật khó khăn bạn từng gặp và cách bạn giải quyết nó.',
    ]),
    ('Phần 2 — Kỹ Năng Mềm', '16A34A', [
        'Kể về một lần bạn phải làm việc với người có quan điểm hoàn toàn khác bạn. Bạn xử lý như thế nào?',
        'Bạn sẽ làm gì khi nhận được nhiều công việc cùng lúc với deadline gấp?',
    ]),
    ('Phần 3 — Kinh Nghiệm Thực Tế', 'B45A09', [
        'Hãy mô tả một dự án bạn tự hào nhất: bạn đóng vai trò gì và kết quả cụ thể là gì?',
        'Kể về một lần bạn mắc sai lầm trong công việc. Bài học bạn rút ra là gì?',
    ]),
]

for part_title, color_hex, questions in sample_parts:
    p = doc.add_paragraph()
    r = p.add_run(part_title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(color_hex)
    for i, q in enumerate(questions, 1):
        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Cm(1.0)
        p2.paragraph_format.space_after = Pt(3)
        p2.add_run(f'Câu hỏi {i}:  ').bold = True
        p2.add_run(f'"{q}"').italic = True

# ══════════════════════════════════════════════════════════════
# 4. CHẤM ĐIỂM CV
# ══════════════════════════════════════════════════════════════
heading(doc, '4. Hệ Thống Chấm Điểm CV', 1, '0051D5')

body(doc,
    'Khi ứng viên nộp hồ sơ, hệ thống AI sẽ đọc CV và chấm điểm theo 10 tiêu chí '
    'được thiết kế sẵn. Điểm tối đa là 10 điểm, ứng viên cần đạt từ 7.5 điểm trở lên '
    'để được mời vào vòng phỏng vấn.'
)

heading(doc, '4.1 Tổng quan các nhóm tiêu chí', 2, '374151')

groups = [
    ('💼', 'Kinh Nghiệm Làm Việc', '3.5 điểm',
     'Số năm kinh nghiệm, mức độ liên quan với vị trí, và có thành tích cụ thể hay không.',
     'EFF6FF'),
    ('🎓', 'Học Vấn & Chứng Chỉ', '2.5 điểm',
     'Bằng cấp, chuyên ngành học, và các chứng chỉ nghề nghiệp liên quan.',
     'F0FDF4'),
    ('⚙', 'Kỹ Năng Kỹ Thuật', '1.0 điểm',
     'Mức độ phù hợp giữa các kỹ năng trong CV với yêu cầu của vị trí tuyển dụng.',
     'FFF7ED'),
    ('🌟', 'Kỹ Năng Mềm & Hồ Sơ', '3.0 điểm',
     'Chất lượng trình bày CV, các dự án thực tế có kết quả đo lường được, '
     'và bằng chứng về năng lực lãnh đạo hoặc làm việc nhóm.',
     'F5F3FF'),
]

t2 = doc.add_table(rows=1, cols=4)
t2.style = 'Table Grid'
for i, h in enumerate(['', 'Nhóm Tiêu Chí', 'Điểm', 'Nội dung xem xét']):
    t2.rows[0].cells[i].text = h
    t2.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    set_cell_bg(t2.rows[0].cells[i], '131B2E')
    t2.rows[0].cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF,0xFF,0xFF)

for icon, name, pts, desc, bg in groups:
    row = t2.add_row()
    row.cells[0].text = icon
    row.cells[1].text = name
    row.cells[1].paragraphs[0].runs[0].bold = True
    row.cells[2].text = pts
    row.cells[2].paragraphs[0].runs[0].bold = True
    row.cells[3].text = desc
    for c in row.cells: set_cell_bg(c, bg)

doc.add_paragraph()
body(doc,
    '⚠  Lưu ý: Điểm số chỉ được HR xem nội bộ. '
    'Ứng viên không nhận được thông tin về điểm số, '
    'dù đạt hay không đạt vòng sàng lọc CV.'
)

# ══════════════════════════════════════════════════════════════
# 5. DASHBOARD QUẢN LÝ
# ══════════════════════════════════════════════════════════════
heading(doc, '5. Dashboard Quản Lý Dành Cho HR', 1, '0051D5')

body(doc,
    'HR đăng nhập vào hệ thống bằng Admin Key để truy cập toàn bộ dữ liệu tuyển dụng. '
    'Giao diện được thiết kế trực quan, không cần đào tạo kỹ thuật.'
)

heading(doc, '5.1 Những gì HR có thể làm', 2, '374151')

hr_features = [
    ('📊', 'Xem tổng quan tuyển dụng',
     'Dashboard hiển thị số lượng CV nhận được, tỷ lệ đạt/không đạt, '
     'số hồ sơ đang chờ xử lý và điểm trung bình của ứng viên.'),
    ('📝', 'Xem chi tiết từng hồ sơ',
     'HR xem điểm số chi tiết theo từng tiêu chí, đọc nhận xét tổng quan của AI, '
     'và tải về CV gốc của ứng viên.'),
    ('🎙', 'Nghe lại phỏng vấn',
     'HR nghe lại 6 câu trả lời ghi âm của ứng viên, chấm điểm từng câu '
     'và ghi chú nhận xét trực tiếp trên hệ thống.'),
    ('🔍', 'Lọc & tìm kiếm',
     'Lọc hồ sơ theo trạng thái (đạt / không đạt / đang chờ), '
     'theo vị trí tuyển dụng, hoặc tìm kiếm theo tên ứng viên.'),
    ('👥', 'Phân quyền truy cập',
     'Admin xem được toàn bộ tính năng. Người dùng khách chỉ xem được '
     'danh sách việc làm công khai, không thể truy cập thông tin ứng viên.'),
]

for icon, title, desc in hr_features:
    bullet(doc, icon, title, desc)

doc.add_paragraph()
heading(doc, '5.2 Thông báo email tự động', 2, '374151')

emails = [
    ('✅ Email mời phỏng vấn',
     'Gửi đến ứng viên đạt ≥ 7.5 điểm. Nội dung lịch sự, chuyên nghiệp, '
     'kèm nút bấm trực tiếp vào trang phỏng vấn. '
     'Không tiết lộ điểm số cho ứng viên.'),
    ('❌ Email thông báo không đạt',
     'Gửi đến ứng viên dưới ngưỡng. Nội dung cảm ơn và khuyến khích '
     'ứng tuyển các vị trí khác phù hợp hơn trong tương lai.'),
]

for title, desc in emails:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.add_run(title + '\n').bold = True
    r = p.add_run(desc)
    r.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

# ══════════════════════════════════════════════════════════════
# 6. CÂU HỎI THƯỜNG GẶP
# ══════════════════════════════════════════════════════════════
heading(doc, '6. Câu Hỏi Thường Gặp', 1, '0051D5')

faqs = [
    ('Ứng viên cần cài phần mềm gì không?',
     'Không. Ứng viên chỉ cần trình duyệt web (Chrome, Safari…) và microphone. '
     'Toàn bộ phỏng vấn diễn ra ngay trên trình duyệt.'),
    ('AI có chấm điểm chính xác không?',
     'AI chấm theo rubric cố định với 10 tiêu chí rõ ràng, cho kết quả nhất quán hơn '
     'việc đọc CV thủ công. Tuy nhiên HR luôn có thể xem lại và điều chỉnh quyết định cuối cùng.'),
    ('Ứng viên có biết điểm số của mình không?',
     'Không. Điểm số chỉ hiển thị nội bộ cho HR. '
     'Ứng viên chỉ nhận email đạt hoặc không đạt, không có thông tin điểm.'),
    ('Mất bao lâu từ lúc nộp hồ sơ đến khi nhận email?',
     'Thường 2–5 phút sau khi nộp. Hệ thống xử lý hoàn toàn tự động, '
     '24/7, không phụ thuộc giờ làm việc của HR.'),
    ('HR có thể thêm vị trí tuyển dụng mới không?',
     'Có. HR quản lý danh sách vị trí trong hệ thống, ứng viên sẽ thấy ngay '
     'khi truy cập trang tuyển dụng.'),
    ('Dữ liệu hồ sơ được lưu ở đâu?',
     'Toàn bộ CV và câu trả lời phỏng vấn được lưu trên máy chủ nội bộ của công ty, '
     'không chia sẻ với bên thứ ba.'),
]

for q, a in faqs:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.add_run('❓  ' + q).bold = True
    p2 = doc.add_paragraph(a)
    p2.paragraph_format.left_indent = Cm(1.2)
    p2.paragraph_format.space_after = Pt(4)
    for run in p2.runs:
        run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

# ── Footer ──────────────────────────────────────────────────────
doc.add_paragraph()
foot = doc.add_paragraph()
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_f = foot.add_run('PAI HR  ·  Tài liệu nội bộ  ·  Tháng 6/2026')
r_f.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
r_f.font.size = Pt(10)

out = '/Users/_qh.fol_/ats_phongvan/PAI_HR_Gioi_Thieu.docx'
doc.save(out)
print(f'Saved: {out}')
